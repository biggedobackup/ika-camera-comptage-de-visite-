"""Exports PDF / Excel : type de contenu, toutes les lignes filtrées, neutralisation des formules Excel."""

import io
import re

import httpx
import pytest
from openpyxl import load_workbook

from app.core.exports import TYPE_EXCEL, TYPE_PDF, generer_excel, neutraliser_formule
from app.utilisateur.model import Role, User

LIGNE_ENTETE = 5  # titre, date, filtres, ligne vide, puis l'en-tête du tableau


def lire_classeur(reponse: httpx.Response) -> list[tuple]:
    feuille = load_workbook(io.BytesIO(reponse.content)).active
    return list(feuille.iter_rows(values_only=True))


def verifier_telechargement(reponse: httpx.Response, type_media: str, nom: str, extension: str) -> None:
    assert reponse.status_code == 200
    assert reponse.headers["content-type"] == type_media
    disposition = reponse.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert re.search(rf'filename="{nom}_\d{{8}}_\d{{4}}\.{extension}"', disposition)
    assert reponse.headers["cache-control"] == "no-store"


# ---------------------------------------------------------------------------
# Neutralisation des formules (injection Excel)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("valeur", ["=1+1", "+33 6 00", "-2+3", "@SOMME(A1)", "\t=1", "\r=1"])
def test_neutraliser_formule(valeur: str) -> None:
    assert neutraliser_formule(valeur) == "'" + valeur


@pytest.mark.parametrize("valeur", ["Alice", "1+1", "a=b", 42, None])
def test_valeurs_sans_danger_inchangees(valeur: object) -> None:
    assert neutraliser_formule(valeur) == valeur


def test_generer_excel_neutralise_toutes_les_cellules_texte() -> None:
    contenu = generer_excel("=Titre", ["=Entete"], [("=HYPERLINK(\"http://x\")",), (12,)], ["=filtre"])
    feuille = load_workbook(io.BytesIO(contenu)).active
    textes = [c for ligne in feuille.iter_rows(values_only=True) for c in ligne if isinstance(c, str)]
    assert all(not t.startswith(("=", "+", "-", "@")) for t in textes)
    assert feuille.cell(row=LIGNE_ENTETE + 2, column=2).value == 12  # nombres conservés


# ---------------------------------------------------------------------------
# Utilisateurs
# ---------------------------------------------------------------------------


async def test_export_pdf_utilisateurs(client_admin: httpx.AsyncClient) -> None:
    reponse = await client_admin.get("/utilisateurs/export/pdf")
    verifier_telechargement(reponse, TYPE_PDF, "utilisateurs", "pdf")
    assert reponse.content.startswith(b"%PDF")


async def test_export_excel_utilisateurs_toutes_les_lignes_filtrees(
    client_admin: httpx.AsyncClient, creer_compte
) -> None:
    for numero in range(1, 26):  # plus d'une page de liste
        await creer_compte(f"export{numero:02d}@ikademo.com", Role.UTILISATEUR, nom_complet=f"Export {numero:02d}")
    await creer_compte("autre@ikademo.com", Role.MANAGER, nom_complet="Autre Personne")

    reponse = await client_admin.get("/utilisateurs/export/excel?q=export&tri=nom_complet&ordre=asc")
    verifier_telechargement(reponse, TYPE_EXCEL, "utilisateurs", "xlsx")
    lignes = lire_classeur(reponse)
    assert lignes[0][0] == "Liste des utilisateurs"
    assert "Recherche : « export »" in lignes[2][0] and "Tri : Nom complet (croissant)" in lignes[2][0]
    assert lignes[LIGNE_ENTETE - 1][:3] == ("N°", "Nom complet", "E-mail")
    donnees = lignes[LIGNE_ENTETE:]
    assert len(donnees) == 25  # toutes les lignes, pas seulement la première page
    assert [ligne[0] for ligne in donnees] == list(range(1, 26))
    assert [ligne[1] for ligne in donnees] == [f"Export {n:02d}" for n in range(1, 26)]


async def test_export_excel_neutralise_les_formules(client_admin: httpx.AsyncClient, creer_compte) -> None:
    await creer_compte("formule@ikademo.com", nom_complet='=HYPERLINK("http://attaquant.example","clic")',
                       telephone="+226 70 00 00 00")
    lignes = lire_classeur(await client_admin.get("/utilisateurs/export/excel?q=formule"))
    ligne = lignes[LIGNE_ENTETE]
    assert ligne[1] == "'" + '=HYPERLINK("http://attaquant.example","clic")'
    assert ligne[3] == "'+226 70 00 00 00"


# ---------------------------------------------------------------------------
# Historique
# ---------------------------------------------------------------------------


async def test_exports_historique(client_admin: httpx.AsyncClient, admin: User) -> None:
    pdf = await client_admin.get("/historique/export/pdf?action=CONNEXION")
    verifier_telechargement(pdf, TYPE_PDF, "historique", "pdf")
    assert pdf.content.startswith(b"%PDF")

    excel = await client_admin.get("/historique/export/excel?action=CONNEXION")
    verifier_telechargement(excel, TYPE_EXCEL, "historique", "xlsx")
    lignes = lire_classeur(excel)
    assert lignes[0][0] == "Historique des actions"
    assert "Action : Connexion" in lignes[2][0]
    donnees = lignes[LIGNE_ENTETE:]
    assert len(donnees) == 1 and donnees[0][3] == "Connexion" and donnees[0][2] == admin.email


async def test_export_refuse_au_role_utilisateur(client_utilisateur: httpx.AsyncClient) -> None:
    for chemin in ("/utilisateurs/export/pdf", "/utilisateurs/export/excel", "/historique/export/pdf",
                   "/historique/export/excel"):
        assert (await client_utilisateur.get(chemin)).status_code == 403, chemin
