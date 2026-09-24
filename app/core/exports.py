"""Exports communs : Excel (openpyxl, anti-injection de formules) et PDF (reportlab, paysage)."""

import enum
import io
from collections.abc import Iterable, Sequence
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any
from urllib.parse import quote
from xml.sax.saxutils import escape

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import settings

BLEU = "1270B8"
CARACTERES_FORMULE = ("=", "+", "-", "@", "\t", "\r")

TYPE_EXCEL = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TYPE_PDF = "application/pdf"


# ---------------------------------------------------------------------------
# Mise en forme des valeurs
# ---------------------------------------------------------------------------


def _horodatage() -> datetime:
    return datetime.now(UTC)


def formater_valeur(valeur: Any) -> str:
    """Représentation texte lisible : dates « jj/mm/aaaa hh:mm », booléens Oui/Non, énumérations, None → ""."""
    if valeur is None:
        return ""
    if isinstance(valeur, bool):
        return "Oui" if valeur else "Non"
    if isinstance(valeur, datetime):
        valeur = valeur.astimezone(UTC) if valeur.tzinfo else valeur
        return valeur.strftime("%d/%m/%Y %H:%M")
    if isinstance(valeur, date):
        return valeur.strftime("%d/%m/%Y")
    if isinstance(valeur, enum.Enum):
        return str(getattr(valeur, "libelle", valeur.value))
    return str(valeur)


def neutraliser_formule(valeur: Any) -> Any:
    """Anti-injection de formules (CSV/Excel injection) : texte commençant par = + - @ tabulation
    ou retour chariot → préfixé par une apostrophe."""
    if isinstance(valeur, str) and valeur.startswith(CARACTERES_FORMULE):
        return "'" + valeur
    return valeur


def _valeur_excel(valeur: Any) -> Any:
    """Nombres conservés en nombres ; tout le reste converti en texte neutralisé."""
    if isinstance(valeur, (int, float, Decimal)) and not isinstance(valeur, bool):
        return valeur
    return neutraliser_formule(formater_valeur(valeur))


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def generer_excel(
    titre: str,
    entetes: Sequence[str],
    lignes: Iterable[Sequence[Any]],
    filtres: Sequence[str] = (),
    *,
    numeroter: bool = True,
) -> bytes:
    """Classeur .xlsx : titre, date de génération, filtres appliqués, puis le tableau.

    Si `numeroter`, une colonne « N° » (1..n) est ajoutée en tête.
    """
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "".join(c for c in titre if c not in '[]:*?/\\')[:31] or "Export"

    colonnes = (["N°"] if numeroter else []) + list(entetes)
    nb_colonnes = len(colonnes)

    def ecrire_texte(ligne: int, texte: str, police: Font) -> None:
        cellule = feuille.cell(row=ligne, column=1, value=neutraliser_formule(texte))
        cellule.data_type = "s"
        cellule.font = police

    ecrire_texte(1, titre, Font(bold=True, size=14, color=BLEU))
    ecrire_texte(2, f"Généré le {_horodatage():%d/%m/%Y à %H:%M} (UTC) — {settings.APP_NAME}", Font(italic=True))
    ecrire_texte(3, "Filtres : " + (" ; ".join(filtres) if filtres else "aucun"), Font(italic=True))

    ligne_entete = 5
    remplissage = PatternFill(fill_type="solid", start_color=BLEU, end_color=BLEU)
    for indice, entete in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=ligne_entete, column=indice, value=neutraliser_formule(entete))
        cellule.data_type = "s"
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = remplissage
        cellule.alignment = Alignment(vertical="center")

    largeurs = [len(str(c)) for c in colonnes]
    numero_ligne = ligne_entete
    for numero, ligne in enumerate(lignes, start=1):
        numero_ligne += 1
        valeurs = ([numero] if numeroter else []) + [_valeur_excel(v) for v in ligne]
        for indice, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=numero_ligne, column=indice, value=valeur)
            if isinstance(valeur, str):
                cellule.data_type = "s"
            largeurs[indice - 1] = max(largeurs[indice - 1], min(len(str(valeur)), 60))

    for indice, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(indice)].width = largeur + 3
    feuille.freeze_panes = feuille.cell(row=ligne_entete + 1, column=1)
    if numero_ligne > ligne_entete:
        feuille.auto_filter.ref = f"A{ligne_entete}:{get_column_letter(nb_colonnes)}{numero_ligne}"

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _pied_de_page(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#6c757d"))
    canvas.drawString(document.leftMargin, 8 * mm, settings.APP_NAME)
    canvas.drawRightString(document.pagesize[0] - document.rightMargin, 8 * mm, f"Page {document.page}")
    canvas.restoreState()


def generer_pdf(
    titre: str,
    entetes: Sequence[str],
    lignes: Iterable[Sequence[Any]],
    filtres: Sequence[str] = (),
    *,
    numeroter: bool = True,
    largeurs: Sequence[float] | None = None,
) -> bytes:
    """Document PDF A4 paysage : titre, date de génération, filtres appliqués et tableau.

    `largeurs` : proportions relatives des colonnes (hors colonne « N° »), réparties sur la largeur utile.
    Tout le texte est échappé (aucune balise reportlab interprétée).
    """
    tampon = io.BytesIO()
    document = SimpleDocTemplate(
        tampon,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=15 * mm,
        title=titre,
        author=settings.APP_NAME,
    )
    styles = getSampleStyleSheet()
    style_titre = ParagraphStyle("titre", parent=styles["Title"], textColor=colors.HexColor("#" + BLEU),
                                 alignment=0, fontSize=16, spaceAfter=4)
    style_info = ParagraphStyle("info", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#495057"))
    style_cellule = ParagraphStyle("cellule", parent=styles["Normal"], fontSize=8, leading=10)
    style_entete = ParagraphStyle("entete", parent=style_cellule, textColor=colors.white, fontName="Helvetica-Bold")

    elements: list[Any] = [
        Paragraph(escape(titre), style_titre),
        Paragraph(escape(f"Généré le {_horodatage():%d/%m/%Y à %H:%M} (UTC)"), style_info),
        Paragraph(escape("Filtres : " + (" ; ".join(filtres) if filtres else "aucun")), style_info),
        Spacer(1, 5 * mm),
    ]

    colonnes = (["N°"] if numeroter else []) + list(entetes)
    donnees: list[list[Any]] = [[Paragraph(escape(str(c)), style_entete) for c in colonnes]]
    for numero, ligne in enumerate(lignes, start=1):
        valeurs = ([str(numero)] if numeroter else []) + [formater_valeur(v) for v in ligne]
        donnees.append([Paragraph(escape(v), style_cellule) for v in valeurs])

    largeur_utile = document.width
    largeur_numero = 12 * mm if numeroter else 0
    proportions = list(largeurs) if largeurs and len(largeurs) == len(entetes) else [1.0] * len(entetes)
    somme = sum(proportions) or 1.0
    largeurs_colonnes = ([largeur_numero] if numeroter else []) + [
        (largeur_utile - largeur_numero) * p / somme for p in proportions
    ]

    tableau = Table(donnees, colWidths=largeurs_colonnes, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLEU)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dee2e6")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for index in range(2, len(donnees), 2):
        style.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#f3f7fb")))
    tableau.setStyle(TableStyle(style))
    elements.append(tableau)
    if len(donnees) == 1:
        elements.append(Spacer(1, 4 * mm))
        elements.append(Paragraph("Aucun élément ne correspond aux filtres appliqués.", style_info))

    document.build(elements, onFirstPage=_pied_de_page, onLaterPages=_pied_de_page)
    return tampon.getvalue()


# ---------------------------------------------------------------------------
# Réponses HTTP
# ---------------------------------------------------------------------------


def _nom_fichier(nom_base: str, extension: str) -> str:
    return f"{nom_base}_{_horodatage():%Y%m%d_%H%M}.{extension}"


def _reponse_fichier(contenu: bytes, nom_fichier: str, type_media: str) -> Response:
    nom_ascii = nom_fichier.encode("ascii", "ignore").decode("ascii") or "export"
    disposition = f"attachment; filename=\"{nom_ascii}\"; filename*=UTF-8''{quote(nom_fichier)}"
    return Response(
        content=contenu,
        media_type=type_media,
        headers={"Content-Disposition": disposition, "Cache-Control": "no-store"},
    )


def reponse_excel(contenu: bytes, nom_base: str) -> Response:
    """Réponse de téléchargement .xlsx (nom : <nom_base>_AAAAMMJJ_HHMM.xlsx)."""
    return _reponse_fichier(contenu, _nom_fichier(nom_base, "xlsx"), TYPE_EXCEL)


def reponse_pdf(contenu: bytes, nom_base: str) -> Response:
    """Réponse de téléchargement .pdf (nom : <nom_base>_AAAAMMJJ_HHMM.pdf)."""
    return _reponse_fichier(contenu, _nom_fichier(nom_base, "pdf"), TYPE_PDF)
